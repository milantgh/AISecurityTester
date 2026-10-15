# -*- coding: utf-8 -*-
"""Report export: Excel / Word / HTML."""
import datetime
import os
from .config import APP_NAME

# ============================================================
# Optional dependencies
# ============================================================
try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    from openpyxl.styles import Font as XFont, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


# ============================================================
# Report export
# ============================================================
EXCEL_HEADERS = [
    "用例 ID", "大类型（一级分类）", "二级分类", "三级分类",
    "前置条件", "测试步骤", "预期结果", "风险等级",
    "防御方案", "OWASP 映射", "说明"
]


def export_excel(records, path):
    if not HAS_XLSX:
        raise RuntimeError("openpyxl not installed")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AI安全测试用例(v1.0)"
    header_font = XFont(bold=True, color="FFFFFF", size=11, name="微软雅黑")
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(EXCEL_HEADERS)
    for c in range(1, len(EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    body_align = Alignment(vertical="top", wrap_text=True)
    for r in records:
        steps = "\n".join(f"{i}. {s}" for i, s in enumerate(r["steps"], 1))
        ws.append([r["case_id"], r["cat"], r["sub"], r["sub3"],
                   r["pre"], steps, r["expect"], r["level"],
                   r["defense"], r["owasp"], r["note"]])
        row_idx = ws.max_row
        for c in range(1, len(EXCEL_HEADERS) + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.alignment = body_align
            cell.border = border
            cell.font = XFont(name="微软雅黑", size=10)
    for i, w in enumerate([10, 18, 14, 26, 40, 50, 40, 10, 50, 16, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26
    wb.save(path)
    return path


def export_word(records, path):
    if not HAS_DOCX:
        raise RuntimeError("python-docx not installed")
    doc = Document()
    doc.add_heading(f"{APP_NAME} - AI Security Test Report", 0)
    doc.add_paragraph(f"生成时间：{datetime.datetime.now():%Y-%m-%d %H:%M:%S}")
    doc.add_paragraph(f"用例数量：{len(records)}")
    for r in records:
        doc.add_heading(f"{r['case_id']}  {r['sub3']}", level=1)
        doc.add_paragraph("")
        doc.add_paragraph(f"一级分类：{r['cat']}")
        doc.add_paragraph(f"二级分类：{r['sub']}")
        doc.add_paragraph(f"OWASP 映射：{r['owasp']}")
        doc.add_paragraph(f"风险等级：{r['level']}")
        doc.add_paragraph(f"前置条件：{r['pre']}")
        doc.add_paragraph("测试步骤：")
        for i, s in enumerate(r["steps"], 1):
            doc.add_paragraph(f"{i}. {s}", style="List Bullet")
        doc.add_paragraph(f"预期结果：{r['expect']}")
        doc.add_paragraph(f"防御方案：{r['defense']}")
        doc.add_paragraph(f"说明：{r['note']}")
    doc.save(path)
    return path


def export_html(records, path):
    rows = ""
    for r in records:
        steps = "<br>".join(f"{i}. {s}" for i, s in enumerate(r["steps"], 1))
        rows += (f"<tr><td>{r['case_id']}</td><td>{r['cat']}</td>"
                 f"<td>{r['sub']}</td><td>{r['sub3']}</td>"
                 f"<td>{r['pre']}</td><td>{steps}</td>"
                 f"<td>{r['expect']}</td><td>{r['level']}</td>"
                 f"<td>{r['defense']}</td><td>{r['owasp']}</td>"
                 f"<td>{r['note']}</td></tr>")
    html = f"""<html><head><meta charset="utf-8"><title>AI Security Test Report</title>
    <style>body{{font-family:Microsoft YaHei;padding:20px}}
    table{{border-collapse:collapse;width:100%}}
    th,td{{border:1px solid #999;padding:6px;vertical-align:top;font-size:12px}}
    th{{background:#4472C4;color:#fff}}</style></head><body>
    <h1>{APP_NAME} - AI Security Test Report</h1>
    <p>生成时间：{datetime.datetime.now():%Y-%m-%d %H:%M:%S}　用例数：{len(records)}</p>
    <table><tr>{''.join(f'<th>{h}</th>' for h in EXCEL_HEADERS)}</tr>{rows}</table>
    </body></html>"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


def export_report(records, base_path):
    written = []
    try:
        written.append(export_excel(records, base_path + ".xlsx"))
    except Exception as e:
        print("Excel export failed:", e)
    try:
        written.append(export_word(records, base_path + ".docx"))
    except Exception as e:
        print("Word export failed:", e)
    if not written:
        written.append(export_html(records, base_path + ".html"))
    return written